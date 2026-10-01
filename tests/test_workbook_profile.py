"""Tests des risques de lecture : en-têtes décalés, fusion et données invalides."""

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from src.extract.workbook_profile import (
    coordinate_key,
    find_header_row,
    merged_label,
    numeric_summary,
    read_profile_table,
    table_number,
)


@pytest.mark.parametrize("name,number", [
    ("Table 5. Monday", 5), ("Table. 8 Thursday", 8), ("Summary", None),
])
def test_table_name_variants(name: str, number: int | None) -> None:
    """Les fautes de ponctuation ne doivent pas perdre une feuille."""
    assert table_number(name) == number


def test_merged_fill_does_not_hide_real_missing_values() -> None:
    """Propager le groupe fusionné, mais laisser un manque hors groupe visible."""
    sheet = Workbook().active
    sheet["B12"] = "COMMUNE"
    sheet["B13"] = "Anfa"
    sheet.merge_cells("B13:B15")
    assert find_header_row(sheet) == 12
    assert merged_label(sheet, 14, 2) == "Anfa"
    assert merged_label(sheet, 16, 2) is None


def test_wrong_hour_header_fails_instead_of_shifting_measurements() -> None:
    """Une heure manquante empêche une lecture silencieusement décalée."""
    sheet = Workbook().active
    sheet["B12"] = "COMMUNE"
    for hour in range(24):
        sheet.cell(12, 13 + hour, hour)
        sheet.cell(12, 37 + hour, hour)
    sheet["M12"] = 23
    with pytest.raises(ValueError, match="Heures inattendues"):
        read_profile_table(sheet, 5)


def test_numeric_profile_reports_invalid_values_without_dropping_count() -> None:
    """Distinguer manque, texte invalide, infini et zéro."""
    result = numeric_summary(pd.Series([None, "bad", np.inf, 0, 2]))
    assert result["rows"] == 5
    assert result["missing"] == 1
    assert result["non_numeric"] == 1
    assert result["non_finite"] == 1
    assert result["non_positive"] == 1


def test_coordinate_text_precision_and_invalid_values() -> None:
    """Accepter la précision textuelle de la source et exposer une coordonnée invalide."""
    assert coordinate_key("33.59180443723376", "-7.637821262107074") == (
        33.59180444, -7.63782126,
    )
    assert coordinate_key(None, "-7.6") is None
    assert coordinate_key("nan", "-7.6") is None
