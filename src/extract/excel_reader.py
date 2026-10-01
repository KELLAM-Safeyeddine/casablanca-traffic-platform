"""Extraire les treize feuilles dans un RAW immuable, versionné par SHA-256.

Seules la structure d'en-têtes et les cellules fusionnées sont résolues. Les temps,
distances, indices et TTI restent bruts jusqu'aux phases de validation/transformation.
"""

import hashlib
import json
import logging
import os
import tempfile
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.extract.workbook_profile import SOURCE_SHA256, read_profile_table, table_number

LOGGER = logging.getLogger(__name__)
METADATA_COLUMNS = (
    "_ingested_at", "_source_file", "_sheet", "_source_sha256", "_payload_sha256",
)


def source_digest(source: Path, expected: str | None = SOURCE_SHA256) -> str:
    """Vérifier l'empreinte source avant toute extraction."""
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if expected is not None and digest != expected:
        raise ValueError(f"Source SHA-256 inattendue : {digest}")
    return digest


def payload_digest(frame: pd.DataFrame) -> str:
    """Hacher les colonnes et toutes les valeurs, sans horodatage d'ingestion."""
    content = frame.drop(columns=list(METADATA_COLUMNS), errors="ignore")
    header = json.dumps(list(content.columns), ensure_ascii=False).encode("utf-8")
    values = pd.util.hash_pandas_object(content, index=False).values.tobytes()
    return hashlib.sha256(header + values).hexdigest()


def verify_raw(path: Path, expected_payload: str | None = None) -> pd.DataFrame:
    """Relire un RAW et refuser une partition corrompue ou privée de métadonnées."""
    frame = pd.read_parquet(path)
    if frame.empty or not set(METADATA_COLUMNS).issubset(frame.columns):
        raise ValueError(f"RAW vide ou sans métadonnées : {path}")
    digest = payload_digest(frame)
    if frame["_payload_sha256"].nunique(dropna=False) != 1:
        raise ValueError(f"Empreinte de contenu ambiguë : {path}")
    if frame["_payload_sha256"].iloc[0] != digest or (
        expected_payload is not None and digest != expected_payload
    ):
        raise ValueError(f"RAW altéré ou incompatible : {path}")
    for column in ("_source_file", "_sheet", "_source_sha256", "_ingested_at"):
        if frame[column].isna().any() or frame[column].nunique() != 1:
            raise ValueError(f"Métadonnée {column} invalide : {path}")
    return frame


def write_immutable_parquet(
    frame: pd.DataFrame, target: Path, source_file: str, sheet: str, source_sha256: str,
) -> Path:
    """Publier atomiquement un Parquet sans remplacer une partition existante."""
    digest = payload_digest(frame)
    if frame.empty:
        raise ValueError("Une partition RAW vide ne doit pas être publiée")
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
        enriched = frame.copy()
        enriched["_ingested_at"] = pd.Timestamp.now(tz="UTC")
        enriched["_source_file"] = source_file
        enriched["_sheet"] = sheet
        enriched["_source_sha256"] = source_sha256
        enriched["_payload_sha256"] = digest
        descriptor, temporary_name = tempfile.mkstemp(dir=target.parent, suffix=".tmp")
        os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            enriched.to_parquet(temporary, index=False)
            try:
                # Le lien dur publie le fichier complet et échoue si la cible existe.
                os.link(temporary, target)
            except FileExistsError:
                pass
        finally:
            temporary.unlink(missing_ok=True)
    stored = verify_raw(target, digest)
    if (stored["_source_sha256"].iloc[0], stored["_sheet"].iloc[0],
        stored["_source_file"].iloc[0]) != (source_sha256, sheet, source_file):
        raise ValueError(f"Provenance RAW incompatible : {target}")
    LOGGER.info("RAW vérifié : %s, %d lignes, feuille=%s", target, len(stored), sheet)
    return target


def inventory_source(source: Path, expected_sha256: str | None = SOURCE_SHA256) -> dict[str, Any]:
    """Inventorier toutes les feuilles et leurs en-têtes sans perdre les libellés source."""
    digest = source_digest(source, expected_sha256)
    workbook = load_workbook(source, data_only=False)
    try:
        sheets = []
        for sheet in workbook:
            number = table_number(sheet.title)
            sheets.append({
                "sheet": sheet.title, "table_number": number,
                "rows": sheet.max_row, "columns": sheet.max_column,
                "merged_ranges": sorted(str(area) for area in sheet.merged_cells.ranges),
                "intro_cells": {cell.coordinate: cell.value
                                for row in sheet.iter_rows(max_row=min(12, sheet.max_row))
                                for cell in row if cell.value is not None},
            })
        numbers = [sheet["table_number"] for sheet in sheets if sheet["table_number"] is not None]
        if sorted(numbers) != list(range(12)) or len(sheets) != 13:
            raise ValueError("La source doit contenir Summary et les tables 0 à 11")
        return {"source_file": source.name, "source_sha256": digest, "sheets": sheets}
    finally:
        workbook.close()


def extract_sheet(
    source: Path, raw_root: Path, number: int | None, expected_sha256: str | None = SOURCE_SHA256,
) -> Path:
    """Lire une feuille et conserver lignes Excel, labels bruts et valeurs d'origine."""
    digest = source_digest(source, expected_sha256)
    filename = "summary.parquet" if number is None else f"table_{number:02d}.parquet"
    target = raw_root / f"source_sha256={digest}" / filename
    if target.exists():
        stored = verify_raw(target)
        if (stored["_source_sha256"].iloc[0] != digest
                or stored["_source_file"].iloc[0] != source.name
                or table_number(stored["_sheet"].iloc[0]) != number):
            raise ValueError(f"Provenance RAW incompatible : {target}")
        return target
    workbook = load_workbook(source, data_only=False)
    try:
        matches = [sheet for sheet in workbook if table_number(sheet.title) == number]
        if len(matches) != 1:
            raise ValueError(f"Feuille absente ou ambiguë : table {number}")
        sheet = matches[0]
        if number is None:
            records = [{f"column_{get_column_letter(column)}": cell.value
                        for column, cell in enumerate(row, start=1)}
                       | {"_excel_row": index} for index, row in enumerate(sheet, start=1)]
            frame = pd.DataFrame(records)
            filename = "summary.parquet"
        else:
            frame = read_profile_table(sheet, number)
            frame["_commune_raw"] = [sheet.cell(int(row), 2).value for row in frame["_excel_row"]]
            frame["_zip_raw"] = [sheet.cell(int(row), 3).value for row in frame["_excel_row"]]
            filename = f"table_{number:02d}.parquet"
        target = raw_root / f"source_sha256={digest}" / filename
        result = write_immutable_parquet(frame, target, source.name, sheet.title, digest)
    finally:
        workbook.close()
    if source_digest(source, expected_sha256) != digest:
        raise RuntimeError("La source a changé pendant l'extraction")
    return result


def extract_workbook(source: Path, raw_root: Path) -> list[Path]:
    """Extraire les treize feuilles, avec un fichier par table et un pour Summary."""
    inventory_source(source)
    return [extract_sheet(source, raw_root, number) for number in [None, *range(12)]]
