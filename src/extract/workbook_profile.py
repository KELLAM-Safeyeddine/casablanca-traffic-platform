"""Profilage en lecture seule : structure, valeurs et diagnostics du classeur source.

Les candidats de correction sont exploratoires. Ce module ne charge aucune base
et n'enregistre jamais le classeur. Les valeurs originales restent dans les diagnostics.
"""

import hashlib
import math
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

SOURCE_SHA256 = "4778abffbe3d7791afa58069fc8a98b6e89995174d4c64401d9367221c42d17d"
DAYS = ("monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday")
TRAFFIC_BASE_COLUMNS = (
    "commune", "zip", "origin_index", "origin_coordinates", "origin_lat",
    "origin_lon", "dest_index", "dest_coordinates", "dest_lat", "dest_lon", "distance_raw",
)


def table_number(name: str) -> int | None:
    """Identifier les tables malgré les points et espaces variables des noms d'onglets."""
    match = re.match(r"^Table[ .]*(\d+)\b", name.strip(), flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def find_header_row(sheet: Worksheet) -> int:
    """Trouver la ligne Commune sans supposer le même décalage pour chaque jour."""
    for row in range(1, min(sheet.max_row, 30) + 1):
        value = sheet.cell(row, 2).value
        if isinstance(value, str) and value.strip().casefold() == "commune":
            return row
    raise ValueError(f"En-tête Commune introuvable : {sheet.title}")


def merged_label(sheet: Worksheet, row: int, column: int) -> Any:
    """Propager seulement une cellule fusionnée ; un vrai manque reste manquant."""
    cell = sheet.cell(row, column)
    if cell.value is not None:
        return cell.value
    for area in sheet.merged_cells.ranges:
        if area.min_col == area.max_col == column and area.min_row <= row <= area.max_row:
            return sheet.cell(area.min_row, column).value
    return None


def read_profile_table(sheet: Worksheet, number: int) -> pd.DataFrame:
    """Lire toutes les lignes du corps avec numéros Excel et labels fusionnés résolus."""
    header = find_header_row(sheet)
    if number >= 5:
        columns = [*TRAFFIC_BASE_COLUMNS]
        for start, prefix in ((13, "time"), (37, "tti")):
            hours = [sheet.cell(header, start + hour).value for hour in range(24)]
            if hours != list(range(24)):
                raise ValueError(f"Heures inattendues : {sheet.title}, colonne {start}")
            columns.extend(f"{prefix}_{hour:02d}" for hour in range(24))
    else:
        last_column = max(
            column for column in range(2, sheet.max_column + 1)
            if sheet.cell(header, column).value is not None
        )
        columns = [str(sheet.cell(header, column).value).strip()
                   for column in range(2, last_column + 1)]
    records = []
    for row in range(header + 1, sheet.max_row + 1):
        values = [sheet.cell(row, column).value for column in range(2, len(columns) + 2)]
        values[:2] = [merged_label(sheet, row, column) for column in (2, 3)]
        records.append(dict(zip(columns, values, strict=True)) | {"_excel_row": row})
    return pd.DataFrame(records)


def numeric_summary(series: pd.Series) -> dict[str, Any]:
    """Compter les manques, échecs de conversion et outliers sans les supprimer."""
    numeric = pd.to_numeric(series, errors="coerce")
    finite = numeric[np.isfinite(numeric)]
    result: dict[str, Any] = {
        "rows": len(series), "missing": int(series.isna().sum()),
        "non_numeric": int((series.notna() & numeric.isna()).sum()),
        "non_finite": int((numeric.notna() & ~np.isfinite(numeric)).sum()),
    }
    if finite.empty:
        return result
    q1, median, q3 = finite.quantile([0.25, 0.5, 0.75])
    iqr = q3 - q1
    return result | {
        "min": float(finite.min()), "q1": float(q1), "median": float(median),
        "q3": float(q3), "p95": float(finite.quantile(0.95)), "max": float(finite.max()),
        "non_positive": int((finite <= 0).sum()),
        "iqr_outliers": int(((finite < q1 - 1.5 * iqr) | (finite > q3 + 1.5 * iqr)).sum()),
    }


def coordinate_key(lat: Any, lon: Any) -> tuple[float, float] | None:
    """Comparer des coordonnées à huit décimales, sans masquer une valeur invalide."""
    try:
        latitude, longitude = float(lat), float(lon)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(latitude) or not math.isfinite(longitude):
        return None
    return round(latitude, 8), round(longitude, 8)


def profile_sheet(sheet: Worksheet, number: int | None) -> dict[str, Any]:
    """Décrire toutes les feuilles, y compris le sommaire sans mesures."""
    formulas = [cell.coordinate for row in sheet for cell in row if cell.data_type == "f"]
    errors = [{"cell": cell.coordinate, "value": cell.value}
              for row in sheet for cell in row if cell.data_type == "e"]
    result: dict[str, Any] = {
        "sheet": sheet.title, "table_number": number, "physical_rows": sheet.max_row,
        "physical_columns": sheet.max_column, "merged_ranges": len(sheet.merged_cells.ranges),
        "formula_count": len(formulas), "excel_errors": errors,
    }
    if number is None:
        return result | {"role": "navigation", "data_rows": 0}
    frame = read_profile_table(sheet, number)
    labels = frame.iloc[:, 0]
    header = find_header_row(sheet)
    raw_missing = sum(sheet.cell(row, 2).value is None
                      for row in range(header + 1, sheet.max_row + 1))
    return result | {
        "role": "traffic" if number >= 5 else "dimension_source",
        "header_row": header, "first_data_row": header + 1, "last_data_row": sheet.max_row,
        "data_rows": len(frame), "columns": list(frame.columns),
        "raw_missing_commune": raw_missing, "missing_commune_after_merged_fill":
        int(labels.isna().sum()), "commune_names": sorted(labels.dropna().unique().tolist()),
        "missing_by_column": {column: int(frame[column].isna().sum()) for column in frame},
        "numeric_by_column": {column: numeric_summary(frame[column]) for column in frame
                              if column not in {frame.columns[0], "_excel_row"}},
    }


def build_measurement_diagnostics(
    tables: dict[int, pd.DataFrame], sheet_names: dict[int, str],
) -> pd.DataFrame:
    """Construire une mesure exploratoire par trajet/jour/heure, sans filtrer les lignes."""
    points = tables[0]
    point_lookup = {
        coordinate_key(row.Latitude, row.Longitude): int(row.Index)
        for row in points.itertuples(index=False)
    }
    if None in point_lookup or len(point_lookup) != len(points):
        raise ValueError("Coordonnées de points invalides ou dupliquées")
    frames = []
    for number in range(5, 12):
        source = tables[number]
        base = source[list(TRAFFIC_BASE_COLUMNS) + ["_excel_row"]].copy()
        base["sheet"] = sheet_names[number]
        base["day_of_week"] = number - 4
        base["day"] = DAYS[number - 5]
        base["distance_km_candidate"] = pd.to_numeric(base["distance_raw"], errors="coerce")
        meters = base["distance_km_candidate"] > 100
        base.loc[meters, "distance_km_candidate"] /= 1000
        base["distance_scaled_candidate"] = meters
        for role in ("origin", "dest"):
            base[f"{role}_point_candidate"] = [point_lookup.get(coordinate_key(lat, lon))
                for lat, lon in zip(base[f"{role}_lat"], base[f"{role}_lon"], strict=True)]
        for hour in range(24):
            frame = base.copy()
            frame["hour"] = hour
            frame["travel_time_raw"] = source[f"time_{hour:02d}"]
            frame["tti_provided"] = source[f"tti_{hour:02d}"]
            frames.append(frame)
    measurements = pd.concat(frames, ignore_index=True)
    raw_time = pd.to_numeric(measurements["travel_time_raw"], errors="coerce")
    # Hypothèse exploratoire : séparateur décimal perdu, réservée aux grands entiers.
    scaled = (raw_time >= 1000) & (raw_time % 1 == 0)
    measurements["time_scaled_candidate"] = scaled
    measurements["travel_time_min_candidate"] = raw_time.where(~scaled, raw_time / 1000)
    reference = measurements.groupby(["origin_point_candidate", "dest_point_candidate"])[
        "travel_time_min_candidate"
    ].transform("min")
    measurements["free_flow_reference_min_candidate"] = reference
    measurements["tti_recalculated_candidate"] = (
        measurements["travel_time_min_candidate"] / reference
    )
    provided = pd.to_numeric(measurements["tti_provided"], errors="coerce")
    measurements["tti_delta_candidate"] = measurements["tti_recalculated_candidate"] - provided
    # Indépendant du TTI recalculé : diagnostic de la formule source à 60 km/h.
    measurements["tti_at_60_kmh_candidate"] = (
        measurements["travel_time_min_candidate"] / measurements["distance_km_candidate"]
    )
    measurements["source_tti_formula_matches"] = np.isclose(
        measurements["tti_at_60_kmh_candidate"], provided, atol=0.002, rtol=0,
    )
    measurements["source_tti_missing_or_below_one"] = provided.isna() | (provided < 1)
    measurements["source_tti_above_five"] = provided > 5
    measurements["tti_absolute_gap_above_0_1"] = measurements["tti_delta_candidate"].abs() > 0.1
    measurements["speed_kmh_candidate"] = (
        60 * measurements["distance_km_candidate"] / measurements["travel_time_min_candidate"]
    )
    return measurements


def profile_identifiers(points: pd.DataFrame, measures: pd.DataFrame) -> dict[str, Any]:
    """Mesurer les incohérences des indices et produire un crosswalk auditable."""
    point_ids = set(points["Index"].astype(int))
    mapping: dict[int, set[int]] = defaultdict(set)
    for role in ("origin", "dest"):
        pairs = measures[[f"{role}_index", f"{role}_point_candidate"]].drop_duplicates()
        for raw_id, point_id in pairs.itertuples(index=False, name=None):
            if pd.notna(point_id):
                mapping[int(raw_id)].add(int(point_id))
    mismatches = measures["origin_index"] != measures["origin_point_candidate"]
    missing = measures["origin_point_candidate"].isna() | measures["dest_point_candidate"].isna()
    raw_ids = set(measures["origin_index"].astype(int)) | set(measures["dest_index"].astype(int))
    distinct = measures.drop_duplicates(["day_of_week", "_excel_row"])
    return {
        "coordinate_point_count": len(points), "traffic_point_count": len(raw_ids),
        "raw_indices_absent_from_dim_point": sorted(raw_ids - point_ids),
        "dim_point_indices_unused_as_raw_index": sorted(point_ids - raw_ids),
        "origin_index_mismatch_rows_across_days": int(
            (distinct["origin_index"] != distinct["origin_point_candidate"]).sum()),
        "origin_index_mismatch_measurements": int(mismatches.sum()),
        "unmatched_coordinates_measurements": int(missing.sum()),
        "ambiguous_source_indices": {str(k): sorted(v) for k, v in mapping.items() if len(v) > 1},
        "point_index_crosswalk": {str(k): sorted(v) for k, v in sorted(mapping.items())},
    }


def summarize_measurements(measures: pd.DataFrame) -> dict[str, Any]:
    """Rassembler les preuves chiffrées des règles candidates de la phase 2."""
    routes = measures.drop_duplicates(["day_of_week", "_excel_row"])
    distance_variants = routes.groupby(["origin_point_candidate", "dest_point_candidate"])[
        "distance_km_candidate"
    ].nunique()
    raw_time = pd.to_numeric(measures["travel_time_raw"], errors="coerce")
    source = pd.to_numeric(measures["tti_provided"], errors="coerce")
    max_per_hour = measures.groupby("hour")["tti_provided"].max()
    return {
        "distance_rows_above_100": int(routes["distance_scaled_candidate"].sum()),
        "distance_rows_in_km": int((~routes["distance_scaled_candidate"]).sum()),
        "distance_variants_per_trajectory": {
            str(k): int(v) for k, v in distance_variants.value_counts().sort_index().items()},
        "trajectories_with_variable_distance": int((distance_variants > 1).sum()),
        "distance_raw": numeric_summary(routes["distance_raw"]),
        "distance_km_candidate": numeric_summary(routes["distance_km_candidate"]),
        "time_raw": numeric_summary(raw_time),
        "time_large_integer_candidates": int(measures["time_scaled_candidate"].sum()),
        "time_large_values_without_candidate": int(
            ((raw_time > 100) & ~measures["time_scaled_candidate"]).sum()),
        "time_min_candidate": numeric_summary(measures["travel_time_min_candidate"]),
        "provided_tti": numeric_summary(source),
        "provided_tti_below_one": int((source < 1).sum()),
        "provided_tti_above_five": int((source > 5).sum()),
        "provided_tti_hourly_maxima": {str(k): float(v) for k, v in max_per_hour.items()},
        "reported_maxima_equal_hours_5_to_23": bool(np.allclose(
            max_per_hour.loc[5:23].to_numpy(), np.arange(5, 24))),
        "recalculated_tti_candidate": numeric_summary(measures["tti_recalculated_candidate"]),
        "tti_signed_delta_mean": float(measures["tti_delta_candidate"].mean()),
        "tti_absolute_delta_mean": float(measures["tti_delta_candidate"].abs().mean()),
        "tti_absolute_delta_p95": float(measures["tti_delta_candidate"].abs().quantile(0.95)),
        "tti_absolute_gap_above_0_1": int(measures["tti_absolute_gap_above_0_1"].sum()),
        "source_formula_60_kmh_matches": int(measures["source_tti_formula_matches"].sum()),
        "scaled_time_source_formula_matches": int((measures["time_scaled_candidate"] &
                                                   measures["source_tti_formula_matches"]).sum()),
        "speed_kmh_candidate": numeric_summary(measures["speed_kmh_candidate"]),
        "source_tti_reference_description": "La source ne précise pas de temps libre de référence",
    }


def summarize_dimensions(tables: dict[int, pd.DataFrame]) -> dict[str, Any]:
    """Comparer les clés et attributs urbains sans inventer d'unités ou de noms."""
    communes = tables[0][["Commune", "ZIP code"]].drop_duplicates()
    canonical_pairs = set(communes.itertuples(index=False, name=None))
    mismatches = {}
    for number in range(1, 5):
        pairs = set(tables[number][["Commune", "ZIP code"]].itertuples(index=False, name=None))
        mismatches[str(number)] = sorted(pairs.symmetric_difference(canonical_pairs))
    population = tables[1]["Population size"]
    household = tables[1]["Household"]
    land = tables[4]
    expected_density = population / (land["Region area (m²)"] / 1_000_000)
    delta_density = (tables[1]["Density"] - expected_density).abs()
    points = tables[0]
    return {
        "commune_count": len(communes), "commune_pairs_mismatch_by_table": mismatches,
        "population_fractional_rows": int((population % 1 != 0).sum()),
        "household_fractional_rows": int((household % 1 != 0).sum()),
        "density_unit": "Non déclarée dans le classeur",
        "density_differs_from_population_per_km2_rows": int((delta_density > 1).sum()),
        "density_population_area_absolute_difference": numeric_summary(delta_density),
        "latitude": numeric_summary(points["Latitude"]),
        "longitude": numeric_summary(points["Longitude"]),
        "points_per_commune": points.groupby("Commune").size().to_dict(),
    }


def profile_workbook(source: Path) -> tuple[dict[str, Any], dict[int, pd.DataFrame], pd.DataFrame]:
    """Profiler toutes les feuilles et garantir la conservation binaire du classeur."""
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    if before != SOURCE_SHA256:
        raise ValueError("La source ne correspond pas au SHA-256 enregistré en phase 0")
    workbook = load_workbook(source, data_only=False)
    try:
        summary = [profile_sheet(sheet, table_number(sheet.title)) for sheet in workbook]
        sheet_names = {table_number(sheet.title): sheet.title for sheet in workbook
                       if table_number(sheet.title) is not None}
        if set(sheet_names) != set(range(12)):
            raise ValueError(f"Tables attendues 0 à 11 : {sheet_names}")
        tables = {number: read_profile_table(workbook[name], number)
                  for number, name in sheet_names.items()}
        measures = build_measurement_diagnostics(tables, sheet_names)
        report = {
            "source_file": source.name, "source_sha256": before, "sheet_count": len(summary),
            "sheets": summary, "identifiers": profile_identifiers(tables[0], measures),
            "measurement_count": len(measures), "trajectory_count": len(
                measures[["origin_point_candidate", "dest_point_candidate"]].drop_duplicates()),
            "measurements": summarize_measurements(measures),
            "dimensions": summarize_dimensions(tables),
        }
    finally:
        workbook.close()
    if hashlib.sha256(source.read_bytes()).hexdigest() != before:
        raise RuntimeError("Le classeur a changé pendant le profilage")
    return report, tables, measures
